#!/usr/bin/env python3
"""Build versioned Supabase seed SQL from the user's source JSON and XLSX."""
from __future__ import annotations

import argparse
import json
import re
import unicodedata
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook

VERSION = "operacao-enem-2026-v1"
SOURCE = "Prompt_APP_Operacao_ENEM_2026.md + Operacao_ENEM_2026.pdf + Serie-Super-Aulas-Assaad.xlsx + links enviados pelo usuário"
ESSAY_URLS = [
    "https://www.youtube.com/watch?v=zXGrzS1jK5o",
    "https://www.youtube.com/watch?v=JieTN96XwrU",
    "https://www.youtube.com/watch?v=BqAl370-Ues",
    "https://www.youtube.com/watch?v=1_gXyMjJoKU",
    "https://www.youtube.com/watch?v=5mjRL5ZeS7I",
    "https://www.youtube.com/watch?v=Fei09Xvy3yI",
    "https://www.youtube.com/watch?v=1VxCWJyLbyk",
    "https://www.youtube.com/watch?v=tstulMtWuqc",
    "https://www.youtube.com/watch?v=SIOL2KWBkXs",
    "https://www.youtube.com/watch?v=-xq1SMpanBk",
]


def parse_duration(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, timedelta):
        return int(value.total_seconds())
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, datetime):
        return value.hour * 3600 + value.minute * 60 + value.second
    text = str(value).strip()
    parts = [int(x) for x in text.split(":")]
    if len(parts) == 3:
        return parts[0] * 3600 + parts[1] * 60 + parts[2]
    if len(parts) == 2:
        return parts[0] * 60 + parts[1]
    return None


def date_iso(value: str) -> str:
    return datetime.strptime(value, "%d/%m/%Y").date().isoformat()


def duration_label(seconds: int | None) -> str | None:
    if seconds is None:
        return None
    minutes = (seconds + 30) // 60
    hours, minutes = divmod(minutes, 60)
    return f"{hours}h{minutes:02d}" if hours else f"{minutes} min"


def slugify(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", ascii_value.lower())).strip("-")


def json_sql(value: object) -> str:
    # Dollar quoting avoids manual escaping while still keeping the data inside a client-supplied SQL migration.
    text = json.dumps(value, ensure_ascii=False, indent=2)
    assert "$seed$" not in text
    return f"$seed${text}$seed$::jsonb"


def recordset(sql_name: str, rows: list[dict], columns: list[tuple[str, str]]) -> str:
    schema = ", ".join(f"{name} {kind}" for name, kind in columns)
    return f"jsonb_to_recordset({json_sql(rows)}) as x({schema})"


def build(args: argparse.Namespace) -> tuple[str, ...]:
    source = Path(args.prompt).read_text(encoding="utf-8")
    match = re.search(r"```json\s*([\s\S]*?)```", source)
    if not match:
        raise SystemExit("Não encontrei o bloco JSON de semente no briefing.")
    data = json.loads(match.group(1))

    sheet = load_workbook(args.xlsx, read_only=True, data_only=True).active
    sheet_rows: dict[int, tuple] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[3] is not None:
            sheet_rows[int(row[3]) + 1] = row
    if set(range(1, 33)) - set(sheet_rows):
        raise SystemExit("A planilha não contém as 32 posições esperadas.")

    resources: list[dict] = []
    material_slugs = {
        "Hospital Questiona": "hospital-questiona",
        "Super Aulas Assaad": "super-aulas-assaad",
        "Redação na Prática": "redacao-na-pratica",
        "Padrões de Matemática": "padroes-de-matematica",
        "Palavras-chave e padrões": "palavras-chave-e-padroes",
        "Aula principal de Humanas": "aula-principal-de-humanas",
        "Vídeo complementar": "video-complementar",
        "Apostilas PreparaENEM Digital": "apostilas-preparaenem-digital",
    }
    for order, item in enumerate(data["materiais"], start=1):
        name = item["nome"]
        if name in ("Aula principal de Humanas", "Vídeo complementar"):
            status = "pending"
        else:
            status = "confirmed"
        kind = "material" if item.get("url") is None else (
            "playlist" if "playlist" in item["url"] else ("video" if "youtube" in item["url"] or "youtu.be" in item["url"] else "site")
        )
        resources.append({
            "slug": material_slugs.get(name, slugify(name)), "name": name,
            "url": item.get("url"), "purpose": item["uso"], "resource_type": kind,
            "verification_status": status, "source": SOURCE, "sort_order": order,
        })
    for resource in resources:
        if resource["slug"] == "super-aulas-assaad":
            resource["purpose"] = "28 aulas do cronograma + 4 aulas complementares (29–32) na biblioteca."
    for order, item in enumerate(data.get("referencias", []), start=20):
        resources.append({
            "slug": f"referencia-{slugify(item['nome'])}", "name": item["nome"],
            "url": item["url"], "purpose": "Referência informativa listada no PDF.",
            "resource_type": "reference", "verification_status": "confirmed",
            "source": SOURCE, "sort_order": order,
        })

    lessons: list[dict] = []
    source_lessons = {int(row["n"]): row for row in data["aulasAssaad"]}
    for number in range(1, 33):
        row = sheet_rows[number]
        direct_url = row[5]
        seconds = parse_duration(row[10])
        available = str(row[14] or "").lower() == "available"
        original = source_lessons.get(number)
        if original:
            title = original["titulo"]
            if number == 24 or original.get("pendente"):
                title = "Título pendente — Aula 24"
            subject = original["disciplina"]
            priority = original.get("prioridade")
            scheduled = date_iso(original["diaPrevisto"])
            playlist_url = original.get("urlPlaylist")
            playlist_index = original.get("indiceNaPlaylist", number)
            label = original.get("duracaoTexto") if seconds is not None else None
            pending = bool(original.get("pendente")) or not available
            is_extra = False
        else:
            title = str(row[4])
            subject = "Complementar"
            priority = None
            scheduled = None
            playlist_url = data["materiais"][1].get("url")
            playlist_index = number
            label = duration_label(seconds)
            pending = not available
            is_extra = True
        lessons.append({
            "lesson_number": number, "title": title, "subject": subject, "priority": priority,
            "duration_seconds": seconds, "duration_label": label,
            "scheduled_date": scheduled, "video_url": direct_url,
            "playlist_url": playlist_url, "playlist_index": playlist_index,
            "availability": "private" if not available else "available",
            "is_extra": is_extra,
            "verification_status": "pending" if pending else "confirmed",
            "source": SOURCE,
        })

    essay_videos: list[dict] = []
    for item, direct_url in zip(data["aulasRedacao"], ESSAY_URLS, strict=True):
        duration_match = re.fullmatch(r"(?:(\d+)h)?(\d+)", item["duracaoTexto"])
        duration_minutes = int(duration_match.group(1) or 0) * 60 + int(duration_match.group(2))
        essay_videos.append({
            "video_number": int(item["n"]), "title": item["titulo"],
            "duration_minutes": duration_minutes, "scheduled_date": date_iso(item["diaPrevisto"]),
            "practice_task": item["tarefaPratica"], "video_url": direct_url,
            "playlist_index": int(item["indiceNaPlaylist"]),
            "verification_status": "confirmed", "source": SOURCE,
        })

    days: list[dict] = []
    for item in data["cronogramaDiario"]:
        days.append({
            "study_date": date_iso(item["data"]), "weekday": item["diaSemana"],
            "phase": "main_plan", "day_type": "study", "focus_subject": item["disciplina"],
            "objective": item["objetivo"], "new_question_goal": int(item["metaQuestoesNovas"]),
            "question_range": None, "source_payload": {}, "source": SOURCE,
        })
    for item in data["retaFinal"]:
        kind = item["tipo"]
        if kind == "revisao_final":
            phase, day_type, focus = "final_review", "review", "Revisão final"
            objective = item["roteiro"].replace("; preparar o deslocamento e dormir no horário habitual.", "")
        elif kind == "prova_1":
            phase, day_type, focus = "exam", "exam_1", "1º dia do Enem"
            parts = (item.get("provas", ""), item.get("levar", ""), item.get("horario", ""))
            objective = ". ".join(part.strip().strip(". ") for part in parts if part and part.strip()) + "."
        elif kind == "prova_2":
            phase, day_type, focus = "exam", "exam_2", "2º dia do Enem"
            parts = (item.get("provas", ""), item.get("estrategia", ""), item.get("horario", ""))
            objective = ". ".join(part.strip().strip(". ") for part in parts if part and part.strip()) + "."
        else:
            phase, day_type, focus = "between_exams", "between_exams", item.get("foco", "Revisão")
            objective = item.get("acao", "")
        days.append({
            "study_date": date_iso(item["data"]), "weekday": item["diaSemana"],
            "phase": phase, "day_type": day_type, "focus_subject": focus,
            "objective": objective, "new_question_goal": None,
            "question_range": item.get("questoes"), "source_payload": {}, "source": SOURCE,
        })

    # Tasks preserve the learning actions, not the personal work/course/lunch schedule.
    tasks: list[dict] = []
    for item in data["cronogramaDiario"]:
        day = date_iso(item["data"])
        subject = item["disciplina"]
        count = int(item["manha"]["questoesNovas"])
        tasks.append({
            "study_date": day, "position": 1, "kind": "question", "title": f"Questões de {subject}",
            "description": f"{count} questões novas no Hospital Questiona. Corrija e classifique os erros.",
            "resource_slug": "hospital-questiona", "target_min": count, "target_max": count,
            "verification_status": "confirmed",
        })
        position = 2
        for number in item["trabalho"]["aulas"]:
            lesson = source_lessons.get(int(number))
            lesson_title = "Título pendente" if lesson and lesson.get("pendente") else (lesson["titulo"] if lesson else "")
            tasks.append({
                "study_date": day, "position": position, "kind": "lesson",
                "title": f"Assistir aula {number}", "description": lesson_title,
                "lesson_number": int(number), "verification_status": "pending" if number == 24 else "confirmed",
            })
            position += 1
        video_number = item["redacao"].get("videoN")
        if video_number:
            essay = data["aulasRedacao"][int(video_number) - 1]
            tasks.append({
                "study_date": day, "position": position, "kind": "essay_video",
                "title": f"Assistir vídeo de Redação {video_number}", "description": essay["titulo"],
                "essay_video_number": int(video_number), "verification_status": "confirmed",
            })
            position += 1
        micro = item["redacao"].get("microtreino")
        if micro:
            tasks.append({
                "study_date": day, "position": position, "kind": "microtraining",
                "title": "Microtreino de Redação", "description": micro,
                "verification_status": "confirmed",
            })
            position += 1
        night = item["noite"]
        tasks.append({
            "study_date": day, "position": position, "kind": "practice",
            "title": "Questões e correção", "description": f"{night['questoes']} questões de livros/PDFs + correção; não usar apostilas PreparaENEM.",
            "resource_slug": "padroes-de-matematica" if subject == "Matemática" else None,
            "target_min": 10, "target_max": 15, "verification_status": "confirmed",
        })
        position += 1
        tasks.append({
            "study_date": day, "position": position, "kind": "review",
            "title": "Revisar erros recentes", "description": "Revisar erros de 24–48 h atrás; se não houver, refazer 5 questões marcadas.",
            "verification_status": "confirmed",
        })

    for item in data["retaFinal"]:
        day = date_iso(item["data"])
        kind = item["tipo"]
        question_range = item.get("questoes", "")
        numbers = [int(n) for n in re.findall(r"\d+", question_range)]
        target_min = numbers[0] if numbers else None
        target_max = numbers[1] if len(numbers) > 1 else target_min
        if kind == "revisao_final":
            text = item["roteiro"]
            # Keep the study-specific action; omit personal sleep/travel scheduling.
            text = text.replace("; preparar o deslocamento e dormir no horário habitual.", "")
            text = text.replace("preparar o deslocamento e dormir no horário habitual.", "")
            title, task_kind = "Roteiro de revisão final", "final_review"
        elif kind in ("prova_1", "prova_2"):
            title = "Dia de prova — confira as orientações"
            task_kind = "exam"
            parts = (item.get("provas", ""), item.get("horario", ""), item.get("levar", ""), item.get("estrategia", ""))
            text = ". ".join(part.strip().strip(". ") for part in parts if part and part.strip()) + "."
        else:
            title = f"Revisão de {item.get('foco', 'conteúdo')}"
            task_kind = "between_exams"
            text = item.get("acao", "")
        tasks.append({
            "study_date": day, "position": 1, "kind": task_kind,
            "title": title, "description": text,
            "target_min": target_min, "target_max": target_max,
            "verification_status": "confirmed",
        })

    version_sql = (
        "-- Versioned seed from supplied prompt, PDF, workbook and user-provided video URLs.\n"
        f"insert into public.plan_versions (version, starts_on, ends_on, source, is_active) values ('{VERSION}', '2026-10-13', '2026-11-15', '{SOURCE.replace(chr(39), chr(39)*2)}', true)\n"
        "on conflict (version) do update set starts_on = excluded.starts_on, ends_on = excluded.ends_on, source = excluded.source, is_active = true;\n\n"
    )
    resource_sql = "insert into public.resources (slug, name, url, purpose, resource_type, verification_status, source, sort_order)\nselect x.slug, x.name, x.url, x.purpose, x.resource_type, x.verification_status, x.source, x.sort_order\nfrom " + recordset("resources", resources, [("slug","text"),("name","text"),("url","text"),("purpose","text"),("resource_type","text"),("verification_status","text"),("source","text"),("sort_order","integer")]) + "\non conflict (slug) do update set name = excluded.name, url = excluded.url, purpose = excluded.purpose, resource_type = excluded.resource_type, verification_status = excluded.verification_status, source = excluded.source, sort_order = excluded.sort_order;\n"
    lesson_sql = "insert into public.lessons (plan_version_id, lesson_number, title, subject, priority, duration_seconds, duration_label, scheduled_date, video_url, playlist_url, playlist_index, availability, is_extra, verification_status, source)\nselect p.id, x.lesson_number, x.title, x.subject, x.priority, x.duration_seconds, x.duration_label, x.scheduled_date, x.video_url, x.playlist_url, x.playlist_index, x.availability, x.is_extra, x.verification_status, x.source\nfrom " + recordset("lessons", lessons, [("lesson_number","smallint"),("title","text"),("subject","text"),("priority","text"),("duration_seconds","integer"),("duration_label","text"),("scheduled_date","date"),("video_url","text"),("playlist_url","text"),("playlist_index","smallint"),("availability","text"),("is_extra","boolean"),("verification_status","text"),("source","text")]) + "\ncross join public.plan_versions p where p.version = '" + VERSION + "'\non conflict (plan_version_id, lesson_number) do update set title = excluded.title, subject = excluded.subject, priority = excluded.priority, duration_seconds = excluded.duration_seconds, duration_label = excluded.duration_label, scheduled_date = excluded.scheduled_date, video_url = excluded.video_url, playlist_url = excluded.playlist_url, playlist_index = excluded.playlist_index, availability = excluded.availability, is_extra = excluded.is_extra, verification_status = excluded.verification_status, source = excluded.source;\n"
    essay_sql = "insert into public.essay_videos (plan_version_id, video_number, title, duration_minutes, scheduled_date, practice_task, video_url, playlist_index, verification_status, source)\nselect p.id, x.video_number, x.title, x.duration_minutes, x.scheduled_date, x.practice_task, x.video_url, x.playlist_index, x.verification_status, x.source\nfrom " + recordset("essay_videos", essay_videos, [("video_number","smallint"),("title","text"),("duration_minutes","integer"),("scheduled_date","date"),("practice_task","text"),("video_url","text"),("playlist_index","smallint"),("verification_status","text"),("source","text")]) + "\ncross join public.plan_versions p where p.version = '" + VERSION + "'\non conflict (plan_version_id, video_number) do update set title = excluded.title, duration_minutes = excluded.duration_minutes, scheduled_date = excluded.scheduled_date, practice_task = excluded.practice_task, video_url = excluded.video_url, playlist_index = excluded.playlist_index, verification_status = excluded.verification_status, source = excluded.source;\n"
    schedule_sql = "insert into public.schedule_days (plan_version_id, study_date, weekday, phase, day_type, focus_subject, objective, new_question_goal, question_range, source_payload, source)\nselect p.id, x.study_date, x.weekday, x.phase, x.day_type, x.focus_subject, x.objective, x.new_question_goal, x.question_range, x.source_payload, x.source\nfrom " + recordset("schedule_days", days, [("study_date","date"),("weekday","text"),("phase","text"),("day_type","text"),("focus_subject","text"),("objective","text"),("new_question_goal","integer"),("question_range","text"),("source_payload","jsonb"),("source","text")]) + "\ncross join public.plan_versions p where p.version = '" + VERSION + "'\non conflict (plan_version_id, study_date) do update set weekday = excluded.weekday, phase = excluded.phase, day_type = excluded.day_type, focus_subject = excluded.focus_subject, objective = excluded.objective, new_question_goal = excluded.new_question_goal, question_range = excluded.question_range, source_payload = excluded.source_payload, source = excluded.source;\n"

    task_columns = [("study_date","date"),("position","smallint"),("kind","text"),("title","text"),("description","text"),("resource_slug","text"),("lesson_number","smallint"),("essay_video_number","smallint"),("target_min","integer"),("target_max","integer"),("verification_status","text")]
    task_template = "insert into public.day_tasks (schedule_day_id, position, kind, title, description, resource_id, lesson_id, essay_video_id, target_min, target_max, verification_status)\nselect d.id, x.position, x.kind, x.title, x.description, r.id, l.id, e.id, x.target_min, x.target_max, x.verification_status\nfrom {recordset}\njoin public.schedule_days d on d.study_date = x.study_date\njoin public.plan_versions p on p.id = d.plan_version_id and p.version = '{version}'\nleft join public.resources r on r.slug = x.resource_slug\nleft join public.lessons l on l.plan_version_id = p.id and l.lesson_number = x.lesson_number\nleft join public.essay_videos e on e.plan_version_id = p.id and e.video_number = x.essay_video_number\non conflict (schedule_day_id, position) do update set kind = excluded.kind, title = excluded.title, description = excluded.description, resource_id = excluded.resource_id, lesson_id = excluded.lesson_id, essay_video_id = excluded.essay_video_id, target_min = excluded.target_min, target_max = excluded.target_max, verification_status = excluded.verification_status;\n"

    args.output.mkdir(parents=True, exist_ok=True)
    outputs = []
    fixed = [
        ("202610130002_seed_content.sql", version_sql + resource_sql),
        ("202610130003_seed_lessons.sql", lesson_sql),
        ("202610130004_seed_essay_videos.sql", essay_sql),
        ("202610130005_seed_schedule_days.sql", schedule_sql),
    ]
    for filename, sql in fixed:
        path = args.output / filename
        path.write_text(sql, encoding="utf-8")
        outputs.append(path)

    sorted_dates = sorted(day["study_date"] for day in days)
    for batch_index, start in enumerate(range(0, len(sorted_dates), 7), start=1):
        batch_dates = set(sorted_dates[start:start + 7])
        batch = [task for task in tasks if task["study_date"] in batch_dates]
        rows = recordset("day_tasks", batch, task_columns)
        sql = task_template.format(recordset=rows, version=VERSION)
        path = args.output / f"202610130{5 + batch_index:03d}_seed_tasks_{batch_index:02}.sql"
        path.write_text(sql, encoding="utf-8")
        outputs.append(path)

    for stale in ("202610130003_seed_schedule.sql", "202610130004_seed_tasks.sql"):
        (args.output / stale).unlink(missing_ok=True)
    return tuple(str(path) for path in outputs)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--prompt", default="/home/ubuntu/upload/Prompt_APP_Operacao_ENEM_2026.md")
    parser.add_argument("--xlsx", default="/home/ubuntu/upload/Serie-Super-Aulas-Assaad.xlsx")
    parser.add_argument("--output", type=Path, default=Path("supabase/migrations"))
    args = parser.parse_args()
    for path in build(args):
        print(path)


if __name__ == "__main__":
    main()
